"""Read-only full workbook/cache comparison and exact August forecast-row reconstruction.

No training loader is invoked, no cache is rebuilt, and no model is imported.
The output is a diagnostic row contract; P1/P3 are ex-post evaluation populations.
"""
import argparse
import collections
import csv
import datetime as dt
import gzip
import hashlib
import io
import json
from pathlib import Path
import re
import sys
import zipfile

import numpy as np
import openpyxl
import pandas as pd

ASK = 'ASK_STEP1_BSTORD_PRC'
BID = 'BID_STEP1_BSTORD_PRC'
IDENTIFIERS = ['ISU_CD', 'TICKER', 'ORD_DD', 'TIME_INTERVAL']
GRID = ['{:02d}:{:02d}:00'.format(m // 60, m % 60) for m in range(540, 921, 10)]
KEYS = ['asset_id', 'date', 'origin_time']
POPS = ['P0_legacy', 'P1_valid_cache_endpoints', 'P2_candidate_clock', 'P3_intersection']
EPSILON = 0.0001


def sha(path):
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for block in iter(lambda: stream.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest()


def verify_inputs(root, manifest):
    records = []
    for row in manifest:
        path = root / row['source_path']
        size = path.stat().st_size
        blob = hashlib.sha1(('blob ' + str(size) + '\0').encode())
        digest = hashlib.sha256()
        with path.open('rb') as stream:
            for block in iter(lambda: stream.read(1 << 20), b''):
                blob.update(block); digest.update(block)
        passed = size == int(row['bytes']) and digest.hexdigest() == row['sha256'] and blob.hexdigest() == row['source_git_blob']
        records.append({'source_path': row['source_path'], 'source_commit': row['source_commit'],
                        'bytes': size, 'sha256': digest.hexdigest(), 'git_blob': blob.hexdigest(), 'matches_pinned_manifest': passed})
    return records


def identifier_text(value):
    if value is None or pd.isna(value):
        return ''
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime('%Y%m%d')
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def clock_text(value):
    # Deliberately preserve elapsed durations and off-hours clocks. In particular,
    # timedelta(2 days, 2 hours) is NOT reduced modulo 24h or turned into a session key.
    return '' if value is None or pd.isna(value) else str(value)


def literal_labels(mid):
    result = []
    for t in range(len(mid) - 1):
        current, future = mid[t], mid[t + 1]
        if current > 0:
            relative = (future - current) / current
            result.append(2 if relative > EPSILON else 0 if relative < -EPSILON else 1)
        else:
            result.append(1)
    return np.asarray(result, dtype=np.int8)


def vector_labels(mid):
    current, future = mid[:, :-1], mid[:, 1:]
    with np.errstate(divide='ignore', invalid='ignore'):
        relative = (future - current) / current
    # NumPy scalar float32 comparisons in the pinned loop use a Python float
    # epsilon. Promote only the already-computed float32 return for comparison.
    compared = relative.astype(np.float64)
    return np.where(current > 0, np.where(compared > EPSILON, 2, np.where(compared < -EPSILON, 0, 1)), 1).astype(np.int8)


def key_digest(frame, label=True):
    columns = KEYS + (['endpoint_time', 'split', 'legacy_label'] if label else [])
    digest = hashlib.sha256()
    for row in frame[columns].itertuples(index=False, name=None):
        digest.update(('\t'.join(str(x) for x in row) + '\n').encode())
    return digest.hexdigest()


def workbook_frame(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    headers = list(next(iterator))
    rows = []
    time_types = collections.Counter()
    time_representations = collections.Counter()
    anomalies = []
    clock_index = headers.index('TIME_INTERVAL')
    for number, values in enumerate(iterator, 2):
        row = list(values)
        time = row[clock_index]
        time_types[type(time).__name__] += 1
        clock = clock_text(time)
        time_representations[clock] += 1
        for name in IDENTIFIERS:
            i = headers.index(name)
            row[i] = clock if name == 'TIME_INTERVAL' else identifier_text(row[i])
        if not re.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d:[0-5]\d', clock) or clock == '23:10:00':
            anomalies.append({'workbook_row_number': number, 'date_token': row[headers.index('ORD_DD')],
                              'time_representation': clock, 'excel_python_type': type(time).__name__})
        rows.append(row)
    metadata = {'sheets': book.sheetnames, 'compared_sheet': sheet.title, 'sheet_rows_including_header': sheet.max_row,
                'sheet_columns': sheet.max_column, 'excel_epoch': str(book.epoch), 'time_python_types': dict(time_types),
                'time_representations': dict(sorted(time_representations.items())), 'anomalies': anomalies}
    book.close()
    with zipfile.ZipFile(path) as archive:
        metadata['formula_elements'] = sum(len(re.findall(rb'<(?:\w+:)?f(?:\s|>)', archive.read(name)))
                                           for name in archive.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name))
        metadata['comment_or_external_link_parts'] = [name for name in archive.namelist() if 'comments' in name.lower() or 'externallink' in name.lower()]
    return pd.DataFrame(rows, columns=headers), metadata


def compare_pair(asset, workbook, cache, atol, rtol):
    schema_equal = list(workbook.columns) == list(cache.columns)
    if not schema_equal:
        raise ValueError('Workbook/cache schema mismatch for ' + asset)
    wkeys = list(workbook[IDENTIFIERS].itertuples(index=False, name=None))
    ckeys = list(cache[IDENTIFIERS].itertuples(index=False, name=None))
    missing_w = sum((collections.Counter(ckeys) - collections.Counter(wkeys)).values())
    missing_c = sum((collections.Counter(wkeys) - collections.Counter(ckeys)).values())
    # Pair duplicate full keys by stable occurrence for comparison; reconstruction
    # separately refuses ambiguous on-grid keys, rather than silently deduplicating.
    def occurrence_keys(keys):
        seen = collections.Counter(); output = []
        for key in keys:
            output.append(key + (seen[key],)); seen[key] += 1
        return pd.MultiIndex.from_tuples(output)
    w = workbook.copy(); w.index = occurrence_keys(wkeys)
    ci = occurrence_keys(ckeys)
    matched = ci.isin(w.index)
    w = w.reindex(ci).reset_index(drop=True)
    row_exact = matched.copy(); row_equivalent = matched.copy()
    column_rows, mismatch_rows = [], []
    for col in cache.columns:
        a, b = w[col], cache[col]
        am, bm = a.isna().to_numpy(), b.isna().to_numpy()
        if col in IDENTIFIERS:
            av = a.fillna('').astype(str).to_numpy(); bv = b.fillna('').astype(str).to_numpy()
            equal = (av == bv) & matched
            exact = equal.copy(); f32diff = np.zeros(len(cache), dtype=bool)
            numeric_exact = numeric_outside = 0
        else:
            av = pd.to_numeric(a, errors='raise').to_numpy(dtype=np.float64)
            bv = pd.to_numeric(b, errors='raise').to_numpy(dtype=np.float64)
            both_missing = am & bm
            exact = ((av == bv) | both_missing) & matched
            equal = np.isclose(av, bv, atol=atol, rtol=rtol, equal_nan=True) & matched
            with np.errstate(over='ignore', invalid='ignore'):
                af, bf = av.astype(np.float32), bv.astype(np.float32)
            f32diff = ~((af == bf) | (np.isnan(af) & np.isnan(bf))) & matched
            numeric_exact = int((~exact & ~am & ~bm & matched).sum())
            numeric_outside = int((~equal & ~am & ~bm & matched).sum())
        missing_diff = (am != bm) & matched
        column_rows.append({'asset_id': asset, 'column': col, 'compared_rows': len(cache),
                            'workbook_missing': int(am[matched].sum()), 'cache_missing': int(bm.sum()),
                            'missingness_mismatches': int(missing_diff.sum()), 'exact_unequal_cells': int((~exact & matched).sum()),
                            'numeric_exact_mismatches': numeric_exact, 'numeric_outside_tolerance': numeric_outside,
                            'float32_value_mismatches': int(f32diff.sum()), 'unmatched_cache_rows': int((~matched).sum())})
        for index in np.flatnonzero((~equal | f32diff) & matched):
            mismatch_rows.append({'asset_id': asset, 'cache_data_row': int(index + 1), 'date_token': cache.iloc[index].ORD_DD,
                                  'clock_token': cache.iloc[index].TIME_INTERVAL, 'column': col,
                                  'missingness_mismatch': int(missing_diff[index]), 'outside_tolerance': int(not equal[index]),
                                  'float32_value_mismatch': int(f32diff[index])})
        row_exact &= exact; row_equivalent &= equal
    cache = cache.copy()
    cache['__file_status'] = np.where(~matched, 'cache_only', np.where(row_exact, 'matched_exact', np.where(row_equivalent, 'matched_tolerance', 'conflict')))
    summary = {'asset_id': asset, 'workbook_rows': len(workbook), 'cache_rows': len(cache), 'schema_equal': schema_equal,
               'key_order_equal': wkeys == ckeys, 'cache_keys_without_workbook': missing_w, 'workbook_keys_without_cache': missing_c,
               'cache_key_duplicates': len(ckeys) - len(set(ckeys)), 'workbook_key_duplicates': len(wkeys) - len(set(wkeys)),
               'exact_unequal_cells': sum(x['exact_unequal_cells'] for x in column_rows),
               'numeric_exact_mismatches': sum(x['numeric_exact_mismatches'] for x in column_rows),
               'numeric_outside_tolerance': sum(x['numeric_outside_tolerance'] for x in column_rows),
               'missingness_mismatches': sum(x['missingness_mismatches'] for x in column_rows),
               'float32_value_mismatches': sum(x['float32_value_mismatches'] for x in column_rows),
               'full_row_exact_matches': int(row_exact.sum()), 'full_row_tolerance_matches': int(row_equivalent.sum()),
               'atol': atol, 'rtol': rtol}
    return cache, summary, column_rows, mismatch_rows


def align_quotes(frame, dates):
    copied = frame.copy()
    copied['TIME_INTERVAL'] = copied.TIME_INTERVAL.astype(str).str.strip()
    selected = copied[copied.ORD_DD.isin(dates) & copied.TIME_INTERVAL.isin(GRID)].copy()
    if selected.duplicated(['ORD_DD', 'TIME_INTERVAL']).any():
        raise ValueError('Duplicate on-grid keys would expand the legacy merge; no deduplication performed')
    selected['__row_exists'] = True
    indexed = selected.set_index(['ORD_DD', 'TIME_INTERVAL'])
    skeleton = pd.MultiIndex.from_product([dates, GRID], names=['ORD_DD', 'TIME_INTERVAL'])
    aligned = indexed.reindex(skeleton)
    aligned['__row_exists'] = aligned['__row_exists'].fillna(False).astype(bool)
    raw = aligned[[ASK, BID]].astype(np.float64)
    filled = raw.groupby(level=0, sort=False).ffill().fillna(0)
    a = filled[ASK].to_numpy(dtype=np.float32).reshape(len(dates), 39)
    b = filled[BID].to_numpy(dtype=np.float32).reshape(len(dates), 39)
    mid = (a + b) / 2.0
    labels = vector_labels(mid)
    # Independent literal loop check for every reconstructed asset-day.
    assert np.array_equal(labels, np.stack([literal_labels(row) for row in mid]))
    sources = {}
    for side, col in [('ask1', ASK), ('bid1', BID)]:
        present = raw[col].notna()
        times = pd.Series(np.tile(GRID, len(dates)), index=skeleton).where(present)
        previous = times.groupby(level=0, sort=False).ffill()
        sources[side] = {'method': np.where(present, 'original', np.where(previous.notna(), 'forward_fill', 'zero_fill')),
                         'source_time': previous.fillna('').to_numpy()}
    return aligned, mid, labels, sources


def reconstruct(asset, cache, workbook, dates, split_map):
    aligned, mid, labels, sources = align_quotes(cache, dates)
    wb, wbmid, wblabel, _ = align_quotes(workbook, dates)
    days = len(dates)
    def endpoint(array, future=False):
        array = np.asarray(array).reshape(days, 39)
        return (array[:, 1:] if future else array[:, :-1]).reshape(-1)
    iso_dates = [dt.datetime.strptime(x, '%Y%m%d').strftime('%Y-%m-%d') for x in dates]
    output = pd.DataFrame({'asset_id': asset, 'date': np.repeat(iso_dates, 38),
                           'origin_time': np.tile(GRID[:-1], days), 'endpoint_time': np.tile(GRID[1:], days),
                           'split': np.repeat([split_map[x] for x in dates], 38), 'legacy_label': labels.reshape(-1)})
    output['no_cross_day_target'] = 1
    output['exact_ten_minute_target'] = 1
    output['scheduled_session_regime'] = np.where(output.date < '2016-08-01', 'pre_20160801_extension', 'from_20160801_extension')
    output['exceptional_session_status'] = 'unresolved'
    status = aligned['__file_status'].fillna('absent_in_both').to_numpy(dtype=object)
    status[(~aligned.__row_exists.to_numpy()) & wb.__row_exists.to_numpy()] = 'workbook_only'
    conflict = np.isin(status, ['conflict', 'cache_only', 'workbook_only'])
    history_conflict = np.maximum.accumulate(conflict.reshape(days, 39), axis=1).reshape(-1)
    valid_endpoints, absent_endpoints, legacy_missing = [], [], []
    for prefix, future in [('origin', False), ('endpoint', True)]:
        output[prefix + '_cache_row_exists'] = endpoint(aligned.__row_exists, future).astype(np.int8)
        output[prefix + '_workbook_row_exists'] = endpoint(wb.__row_exists, future).astype(np.int8)
        output[prefix + '_workbook_cache_status'] = endpoint(status, future)
        output[prefix + '_file_value_conflict'] = endpoint(conflict, future).astype(np.int8)
        present, finite, positive, values = {}, {}, {}, {}
        for side, col in [('ask1', ASK), ('bid1', BID)]:
            value = endpoint(aligned[col].to_numpy(dtype=np.float64), future)
            values[side] = value
            present[side] = ~np.isnan(value)
            finite[side] = np.isfinite(value)
            positive[side] = finite[side] & (value > 0)
            for suffix, flag in [('present', present[side]), ('finite', finite[side]), ('finite_positive', positive[side]),
                                 ('nonpositive', finite[side] & (value <= 0))]:
                output[prefix + '_cache_' + side + '_' + suffix] = flag.astype(np.int8)
            output[prefix + '_' + side + '_fill_method'] = endpoint(sources[side]['method'], future)
            output[prefix + '_' + side + '_source_time'] = endpoint(sources[side]['source_time'], future)
            output[prefix + '_workbook_' + side + '_present'] = (~np.isnan(endpoint(wb[col].to_numpy(dtype=np.float64), future))).astype(np.int8)
        both_finite = finite['ask1'] & finite['bid1']
        both_positive = positive['ask1'] & positive['bid1']
        crossed = both_finite & (values['bid1'] > values['ask1'])
        locked = both_positive & (values['bid1'] == values['ask1'])
        output[prefix + '_crossed'] = crossed.astype(np.int8)
        output[prefix + '_locked_positive'] = locked.astype(np.int8)
        valid = both_positive & ~crossed
        output[prefix + '_valid_original_cache_quote'] = valid.astype(np.int8)
        valid_endpoints.append(valid)
        absent_endpoints.append(~(present['ask1'] & present['bid1']))
        # Prior audit treated absent/nonfinite quotes as missing. Keep that
        # definition distinct from actual NaN missingness and strict positivity.
        legacy_missing.append(~both_finite)
    output['origin_history_file_conflict'] = endpoint(history_conflict).astype(np.int8)
    output['workbook_cache_filled_mid_equal_origin'] = (mid[:, :-1] == wbmid[:, :-1]).reshape(-1).astype(np.int8)
    output['workbook_cache_filled_mid_equal_endpoint'] = (mid[:, 1:] == wbmid[:, 1:]).reshape(-1).astype(np.int8)
    output['workbook_cache_label_match'] = (labels == wblabel).reshape(-1).astype(np.int8)
    missing = legacy_missing[0] | legacy_missing[1]
    output['missing_cache_quote_any_endpoint'] = (absent_endpoints[0] | absent_endpoints[1]).astype(np.int8)
    output['historical_missing_cache_endpoint'] = missing.astype(np.int8)
    output['historical_missing_endpoint_kind'] = np.where(legacy_missing[0], np.where(legacy_missing[1], 'both', 'current'), np.where(legacy_missing[1], 'future', 'none'))
    output['equal_filled_mids'] = (mid[:, :-1] == mid[:, 1:]).reshape(-1).astype(np.int8)
    output['historical_missing_equal_filled_flat'] = (missing & (output.equal_filled_mids.to_numpy() == 1) & (labels.reshape(-1) == 1)).astype(np.int8)
    clock = (output.origin_time >= '09:10:00') & (output.endpoint_time <= '14:40:00')
    output['candidate_clock'] = clock.astype(np.int8)
    output['P0_legacy'] = 1
    output['P1_valid_cache_endpoints'] = (valid_endpoints[0] & valid_endpoints[1]).astype(np.int8)
    output['P2_candidate_clock'] = clock.astype(np.int8)
    output['P3_intersection'] = (clock & (output.P1_valid_cache_endpoints == 1)).astype(np.int8)
    # Explicit current peer availability depends on this asset's origin alone.
    output['peer_origin_available_cache'] = output.origin_valid_original_cache_quote
    reasons = []
    for row in output.itertuples(index=False):
        parts = []
        for prefix in ['origin', 'endpoint']:
            if not getattr(row, prefix + '_cache_row_exists'): parts.append(prefix + '_row_absent')
            for side in ['ask1', 'bid1']:
                if not getattr(row, prefix + '_cache_' + side + '_present'): parts.append(prefix + '_' + side + '_missing')
                elif not getattr(row, prefix + '_cache_' + side + '_finite'): parts.append(prefix + '_' + side + '_nonfinite')
                elif not getattr(row, prefix + '_cache_' + side + '_finite_positive'): parts.append(prefix + '_' + side + '_nonpositive')
            if getattr(row, prefix + '_crossed'): parts.append(prefix + '_crossed')
        reasons.append('|'.join(parts) if parts else 'none')
    output['P1_exclusion_reasons'] = reasons
    output['P2_exclusion_reason'] = np.where(clock, 'none', np.where(output.origin_time < '09:10:00', 'origin_before_0910', 'endpoint_after_1440'))
    assert not output.duplicated(KEYS).any()
    assert int(output.P2_candidate_clock.sum()) == days * 33
    assert np.array_equal(output.P3_intersection.to_numpy(), (output.P1_valid_cache_endpoints & output.P2_candidate_clock).to_numpy())
    assert np.array_equal(output.peer_origin_available_cache, output.origin_valid_original_cache_quote)
    return output


def population_counts(frame, group_columns):
    output = []
    grouper = group_columns[0] if len(group_columns) == 1 else group_columns
    for key, group in frame.groupby(grouper, sort=True):
        key = (key,) if len(group_columns) == 1 else key
        identity = dict(zip(group_columns, key))
        for population in POPS:
            retained = group[group[population] == 1]
            classes = retained.legacy_label.value_counts()
            row = dict(identity, population=population, legacy_rows=len(group), retained=len(retained), excluded=len(group)-len(retained))
            for label, name in [(0, 'down'), (1, 'flat'), (2, 'up')]:
                count = int(classes.get(label, 0)); row[name + '_count'] = count
                row[name + '_proportion'] = count / len(retained) if len(retained) else None
            output.append(row)
    return output


def deterministic_gzip(path):
    return io.TextIOWrapper(gzip.GzipFile(filename='', mode='wb', fileobj=path.open('xb'), mtime=0), encoding='utf-8', newline='')


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--inputs', required=True, type=Path)
    ap.add_argument('--repository', required=True, type=Path)
    ap.add_argument('--atol', type=float, default=1e-12)
    ap.add_argument('--rtol', type=float, default=1e-12)
    args = ap.parse_args()
    root = args.repository
    if (root / 'artifacts/qf_data_review/20261001/forecast_rows.csv.gz').exists():
        raise FileExistsError('Existing diagnostic rows preserved; use a new audit destination.')
    docs = root / 'docs/qf_data_review/20261001'; docs.mkdir(parents=True, exist_ok=True)
    artifacts = root / 'artifacts/qf_data_review/20261001'; artifacts.mkdir(parents=True, exist_ok=True)
    frozen_index = docs / 'prediction_index.json'
    assert frozen_index.exists(), 'Freeze the prediction index before constructing populations.'
    freeze_hash = sha(frozen_index)
    manifest_path = root / 'recovery/input_manifest.csv'
    manifest = list(csv.DictReader(manifest_path.open()))
    assert len(manifest) == 54
    before = verify_inputs(args.inputs, manifest)
    pd.DataFrame(before).to_csv(docs / 'input_verification_before.csv', index=False)
    assert all(row['matches_pinned_manifest'] for row in before)
    stems = sorted(Path(row['source_path']).stem for row in manifest if row['category'] == 'existing_runtime_csv_cache')
    workbook_stems = sorted(Path(row['source_path']).stem for row in manifest if row['category'] == 'supplied_aggregate_workbook')
    assert stems == workbook_stems and len(stems) == 27
    date_sets = [set(pd.read_csv(args.inputs / 'data/processed' / (stem + '.csv'), usecols=['ORD_DD'], dtype=str).ORD_DD) for stem in stems]
    common = sorted(set.intersection(*date_sets)); union = sorted(set.union(*date_sets))
    train_end, val_end = int(len(common) * .7), int(len(common) * .8)
    split_map = {day: 'train' if i < train_end else 'validation' if i < val_end else 'development_holdout' for i, day in enumerate(common)}
    pd.DataFrame({'date': common, 'split': [split_map[d] for d in common]}).to_csv(docs / 'split_dates.csv', index=False)
    expected = {'train': 353970, 'validation': 50274, 'development_holdout': 101574}
    all_rows, comparisons, column_results, mismatches, metadata, clock_rows = [], [], [], [], [], []
    errors = []
    for asset in stems:
        try:
            workbook_path = args.inputs / 'data/raw/KRXaggData' / (asset + '.xlsx')
            cache_path = args.inputs / 'data/processed' / (asset + '.csv')
            workbook, meta = workbook_frame(workbook_path)
            cache = pd.read_csv(cache_path, dtype={name: str for name in IDENTIFIERS})
            for name in IDENTIFIERS:
                cache[name] = cache[name].fillna('')
            compared, summary, columns, issues = compare_pair(asset, workbook, cache, args.atol, args.rtol)
            comparisons.append(summary); column_results.extend(columns); mismatches.extend(issues)
            meta['asset_id'] = asset; metadata.append(meta)
            for date, group in cache.groupby('ORD_DD', sort=True):
                times = group.TIME_INTERVAL.astype(str).str.strip()
                valid_clock = times.str.fullmatch(r'(?:[01]\d|2[0-3]):[0-5]\d:[0-5]\d')
                on_grid = times.isin(GRID)
                clock_rows.append({'asset_id': asset, 'date': date, 'common_date': int(date in split_map),
                                   'split': split_map.get(date, 'outside_common_dates'), 'cache_rows': len(group),
                                   'on_legacy_grid_rows': int(on_grid.sum()), 'off_legacy_grid_rows': int((~on_grid).sum()),
                                   'invalid_clock_representation_rows': int((~valid_clock).sum()),
                                   'first_normal_clock': times[valid_clock].min() if valid_clock.any() else '',
                                   'last_normal_clock': times[valid_clock].max() if valid_clock.any() else '',
                                   'scheduled_session_regime': 'pre_20160801_extension' if date < '20160801' else 'from_20160801_extension'})
            rows = reconstruct(asset, compared, workbook, common, split_map)
            summary['forecast_rows'] = len(rows)
            summary['workbook_cache_label_mismatches'] = int((rows.workbook_cache_label_match == 0).sum())
            summary['forecast_origin_history_conflicts'] = int(rows.origin_history_file_conflict.sum())
            summary['forecast_file_conflict_any_endpoint'] = int(((rows.origin_file_value_conflict == 1) | (rows.endpoint_file_value_conflict == 1)).sum())
            all_rows.append(rows)
            print(json.dumps({'asset': asset, 'workbook_rows': len(workbook), 'cache_rows': len(cache),
                              'numeric_outside_tolerance': summary['numeric_outside_tolerance'],
                              'missingness_mismatches': summary['missingness_mismatches'],
                              'float32_value_mismatches': summary['float32_value_mismatches'],
                              'legacy_labels': len(rows), 'label_mismatches': summary['workbook_cache_label_mismatches']}), flush=True)
        except Exception as error:
            errors.append({'asset_id': asset, 'error_type': type(error).__name__, 'detail': str(error)})
            print(json.dumps(errors[-1]), flush=True)
    pd.DataFrame(comparisons).to_csv(docs / 'workbook_cache_comparison.csv', index=False)
    pd.DataFrame(column_results).to_csv(docs / 'workbook_cache_columns.csv', index=False)
    pd.DataFrame(clock_rows).to_csv(docs / 'daily_clock_summary.csv', index=False)
    (docs / 'workbook_metadata.json').write_text(json.dumps(metadata, indent=2) + '\n')
    (docs / 'comparison_errors.json').write_text(json.dumps(errors, indent=2) + '\n')
    if mismatches:
        pd.DataFrame(mismatches).to_csv(artifacts / 'workbook_cache_mismatch_locations.csv.gz', index=False, compression={'method': 'gzip', 'mtime': 0})
    if not all_rows:
        raise RuntimeError('No forecast rows reconstructed; see comparison_errors.json')
    frame = pd.concat(all_rows, ignore_index=True).sort_values(KEYS, kind='stable').reset_index(drop=True)
    assert not frame.duplicated(KEYS).any()
    with deterministic_gzip(artifacts / 'forecast_rows.csv.gz') as output:
        frame.to_csv(output, index=False, lineterminator='\n')
    for name, groups in [('population_counts_by_split.csv', ['split']), ('population_counts_by_asset.csv', ['asset_id']),
                         ('population_counts_by_asset_split.csv', ['asset_id', 'split']),
                         ('population_counts_by_regime_split.csv', ['scheduled_session_regime', 'split']),
                         ('population_counts_by_origin_split.csv', ['origin_time', 'split'])]:
        pd.DataFrame(population_counts(frame, groups)).to_csv(docs / name, index=False)
    flag_columns = [c for c in frame if c.endswith(('_present', '_finite', '_finite_positive', '_nonpositive', '_crossed', '_locked_positive', '_cache_row_exists', '_workbook_row_exists'))] + ['missing_cache_quote_any_endpoint', 'historical_missing_cache_endpoint', 'historical_missing_equal_filled_flat']
    missing_rows = []
    for (split, asset), group in frame.groupby(['split', 'asset_id'], sort=True):
        for col in flag_columns:
            missing_rows.append({'split': split, 'asset_id': asset, 'flag': col, 'count_true': int(group[col].sum()), 'denominator': len(group)})
        for col in [c for c in frame if c.endswith(('_fill_method', '_workbook_cache_status'))] + ['historical_missing_endpoint_kind']:
            for value, count in group[col].value_counts().items():
                missing_rows.append({'split': split, 'asset_id': asset, 'flag': col + '=' + str(value), 'count_true': int(count), 'denominator': len(group)})
    pd.DataFrame(missing_rows).to_csv(docs / 'missingness_by_asset_split.csv', index=False)
    after = verify_inputs(args.inputs, manifest)
    pd.DataFrame(after).to_csv(docs / 'input_verification_after.csv', index=False)
    assert before == after and all(row['matches_pinned_manifest'] for row in after)
    assert sha(frozen_index) == freeze_hash
    actual_split = {key: int(value) for key, value in frame.groupby('split').size().items()}
    actual_missing = {key: int(value) for key, value in frame.groupby('split').historical_missing_cache_endpoint.sum().items()}
    checks = {'workbook_pairs_compared': len(comparisons), 'comparison_errors': errors,
              'input_files_verified_before_and_after': len(before), 'input_hashes_unchanged': before == after,
              'common_dates': len(common), 'common_first': common[0], 'common_last': common[-1],
              'union_dates': len(union), 'union_dates_excluded_by_intersection': sorted(set(union)-set(common)),
              'forecast_rows': len(frame), 'split_rows': actual_split, 'expected_split_rows': expected,
              'split_counts_match_reference': actual_split == expected, 'duplicate_keys': int(frame.duplicated(KEYS).sum()),
              'row_key_sha256': key_digest(frame, False), 'row_key_label_sha256': key_digest(frame, True),
              'digest_serialization': 'UTF8 LF-terminated TAB-separated columns asset_id,date,origin_time[,endpoint_time,split,legacy_label], sorted lexicographically by asset_id,date,origin_time',
              'split_digests': {split: key_digest(group) for split, group in frame.groupby('split', sort=True)},
              'historical_missing_endpoint': int(frame.historical_missing_cache_endpoint.sum()),
              'historical_missing_endpoint_by_split': actual_missing,
              'historical_missing_equal_filled_flat': int(frame.historical_missing_equal_filled_flat.sum()),
              'historical_missing_endpoint_kind': {str(k): int(v) for k, v in frame.historical_missing_endpoint_kind.value_counts().items()},
              'missing_reference_matches': int(frame.historical_missing_cache_endpoint.sum()) == 34600 and actual_missing == {'train':28195,'validation':1994,'development_holdout':4411},
              'overlapping_subset_reference_matches': int(frame.historical_missing_equal_filled_flat.sum()) == 27366,
              'workbook_cache_label_mismatches': int((frame.workbook_cache_label_match == 0).sum()),
              'P1_excluded': int((frame.P1_valid_cache_endpoints == 0).sum()),
              'population_totals': {name: int(frame[name].sum()) for name in POPS},
              'prediction_index_sha256_at_start_and_end': freeze_hash,
              'literal_loop_label_equivalence': 'checked every asset-day; exact equality',
              'restricted_population_model_scores_computed': False,
              'python': sys.version.split()[0], 'numpy': np.__version__, 'pandas': pd.__version__, 'openpyxl': openpyxl.__version__}
    checks['data_ready'] = (not errors and len(comparisons) == 27 and checks['split_counts_match_reference']
                            and checks['missing_reference_matches'] and checks['overlapping_subset_reference_matches'])
    (docs / 'data_checks.json').write_text(json.dumps(checks, indent=2) + '\n')
    schema = {col: {'storage_dtype': str(frame[col].dtype)} for col in frame.columns}
    (docs / 'schema_columns.json').write_text(json.dumps(schema, indent=2) + '\n')
    print(json.dumps(checks), flush=True)


if __name__ == '__main__':
    main()
